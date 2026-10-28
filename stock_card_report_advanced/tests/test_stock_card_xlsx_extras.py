# -*- coding: utf-8 -*-
import io
from datetime import timedelta

from odoo import fields
from odoo.tests import tagged

from .common import StockCardCommon


@tagged("post_install", "-at_install", "stock_card")
class TestStockCardXlsxPartnerNotes(StockCardCommon):
    """Detailed report, XLSX only: Partner and Notes columns after Balance Value."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        group_manager = cls.env.ref("stock.group_stock_manager")
        cls.picker = cls.env["res.users"].create({
            "name": "SC Picker", "login": "sc_picker",
            "group_ids": [(6, 0, [cls.env.ref("stock.group_stock_user").id])],
        })
        cls.counter = cls.env["res.users"].create({
            "name": "SC Counter", "login": "sc_counter",
            "group_ids": [(6, 0, [group_manager.id])],
        })
        cls.scrapper = cls.env["res.users"].create({
            "name": "SC Scrapper", "login": "sc_scrapper",
            "group_ids": [(6, 0, [group_manager.id])],
        })
        cls.vendor = cls.env["res.partner"].create({"name": "SC Vendor"})
        # a user whose contact belongs to the company, as in most databases
        cls.picker.partner_id.parent_id = cls.env["res.partner"].create({"name": "SC Company", "is_company": True})

    def _today_options(self, **kw):
        today = fields.Date.today()
        return self._options(date_from=fields.Date.to_string(today - timedelta(days=1)),
                             date_to=fields.Date.to_string(today + timedelta(days=1)), **kw)

    def _receipt(self, qty):
        picking = self.env["stock.picking"].create({
            "picking_type_id": self.warehouse.in_type_id.id,
            "location_id": self.supplier.id,
            "location_dest_id": self.stock.id,
            "partner_id": self.vendor.id,
            "user_id": self.picker.id,
            "note": "<p>Deliver to <b>dock 2</b></p><p>Fragile</p>",
            "move_ids": [(0, 0, {
                "product_id": self.product.id, "product_uom_qty": qty,
                "product_uom": self.product.uom_id.id,
                "location_id": self.supplier.id, "location_dest_id": self.stock.id,
            })],
        })
        picking.action_confirm()
        picking.move_ids.write({"quantity": qty, "picked": True})
        picking.button_validate()
        self.assertEqual(picking.state, "done")
        return picking

    def _apply_count(self, user, quantity):
        quant = self.env["stock.quant"].with_user(user).with_context(inventory_mode=True).create({
            "product_id": self.product.id,
            "location_id": self.stock.id,
            "inventory_quantity": quantity,
        })
        quant.action_apply_inventory()

    def test_01_partner_and_notes_per_move_type(self):
        self._receipt(10)
        self._apply_count(self.counter, 13)   # +3: adjustment increase
        self._apply_count(self.counter, 11)   # -2: adjustment decrease
        scrap = self.env["stock.scrap"].with_user(self.scrapper).create({
            "product_id": self.product.id,
            "product_uom_id": self.product.uom_id.id,
            "scrap_qty": 1,
            "location_id": self.stock.id,
        })
        scrap.action_validate()

        data = self.engine.get_report_data(self._today_options(), xlsx_extras=True)
        by_type = {l["move_type"]: l for l in data["groups"][0]["lines"]}
        self.assertEqual(set(by_type), {"receipt", "inventory_in", "inventory_out", "scrap"})
        # transfer -> its responsible; notes -> plain text of the transfer note
        self.assertEqual(by_type["receipt"]["partner_name"], self.picker.name)  # not "SC Company, SC Picker"
        self.assertIn("dock 2", by_type["receipt"]["picking_note"])
        self.assertIn("Fragile", by_type["receipt"]["picking_note"])
        self.assertNotIn("<p>", by_type["receipt"]["picking_note"])
        # adjustments -> user who applied them; scrap -> user who created it
        self.assertEqual(by_type["inventory_in"]["partner_name"], self.counter.name)
        self.assertEqual(by_type["inventory_out"]["partner_name"], self.counter.name)
        self.assertEqual(by_type["scrap"]["partner_name"], self.scrapper.name)
        self.assertEqual(by_type["inventory_in"]["picking_note"], "")

        # XLSX: Partner and Notes right after Balance Value
        _name, fileobj, _size = self.env["stock.card.report.xlsx"].generate_xlsx(self._today_options())
        try:
            import openpyxl as xl
        except ImportError:
            fileobj.close()
            return
        with fileobj:
            sheet = xl.load_workbook(io.BytesIO(fileobj.read()))["Stock Card"]
        rows = list(sheet.iter_rows(values_only=True))
        header = next(r for r in rows if r[0] == "Date")
        names = [c for c in header if c]
        self.assertEqual(names[-3:], [names[-3], "Partner", "Notes"])
        self.assertTrue(names[-3].startswith("Balance Value"))
        flat = [str(c) for r in rows for c in r if c]
        self.assertIn(self.picker.name, flat)
        self.assertIn(self.counter.name, flat)

    def test_02_not_in_pdf_web_or_totals(self):
        self._receipt(4)
        # web / PDF payload: extras are not even queried
        line = self.engine.get_report_data(self._today_options())["groups"][0]["lines"][0]
        self.assertNotIn("partner_name", line)
        self.assertNotIn("picking_note", line)
        rpc = self.engine.rpc_get_lines(self._today_options())
        self.assertNotIn("partner_name", next(iter(rpc.values()))[0])
        html, _fmt = self.env["ir.actions.report"]._render_qweb_html(
            "stock_card_report_advanced.report_stock_card", [], data={"options": self._today_options()})
        html = html.decode() if isinstance(html, bytes) else html
        self.assertNotIn(">Partner<", html)
        self.assertNotIn("dock 2", html)
        # totals report: no partner / notes columns in the XLSX
        _name, fileobj, _size = self.env["stock.card.report.xlsx"].generate_xlsx(
            self._today_options(report_mode="summary"))
        try:
            import openpyxl as xl
        except ImportError:
            fileobj.close()
            return
        with fileobj:
            sheet = xl.load_workbook(io.BytesIO(fileobj.read()))["Stock Card"]
        header = next(r for r in sheet.iter_rows(values_only=True) if r[0] == "Date")
        self.assertNotIn("Partner", header)
        self.assertNotIn("Notes", header)
